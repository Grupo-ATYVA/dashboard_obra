"""Medição PLS (Caixa) — leitura da 'PLS_HIS_MED<nº>...xlsx' e marcações da medição seguinte.

A planilha PLS-Habitação guarda, para cada serviço (linha) e cada unidade/trecho (coluna a partir de BD),
o número da medição em que aquela unidade foi medida (célula vazia = não medida). Os estados vêm de:
  - número na célula ........ medido (anterior, ou no período se for o nº da medição atual)
  - preenchimento vermelho ... glosado
  - hachurado (lightUp) ...... não se aplica ao serviço
  - preenchimento liso vazio . em aberto (pode ser medido)
Células mescladas contam como UM evento (vale a célula âncora).

Só lê as planilhas; as marcações feitas no painel ficam em dados/uau/pls_medicao_estado.json e
podem ser exportadas num XLSX no formato da PLS para colar na planilha oficial.
"""
import datetime
import glob
import json
import os
import re
import threading
import time

CONJUNTOS = {
    "cj03": {"nome": "CJ03 — Regina Célia", "chaves": ("regina",)},
    "cj05": {"nome": "CJ05 — Joaquim Roriz", "chaves": ("joaquim", "roriz")},
}
COR_VERMELHA = "FFFF0000"

_lock_estado = threading.RLock()
_cache_mem = {}                       # cj -> {"mtime":..., "dados":...}
_parse_threads = {}                   # cj -> {"thread":..., "erro":...}
_parse_lock = threading.Lock()


# ───────────────────────── localizar arquivos ─────────────────────────
def localizar(pasta):
    """{'cj03': caminho, 'cj05': caminho} — o PLS mais recente de cada conjunto."""
    achados = {}
    for p in glob.glob(os.path.join(pasta, "*PLS_HIS_MED*.xlsx")):   # aceita prefixo no nome (ex.: "_PLS_HIS_MED11...")
        nome = os.path.basename(p)
        if nome.startswith("~$"):
            continue
        baixo = nome.lower()
        for cj, cfg in CONJUNTOS.items():
            if any(k in baixo for k in cfg["chaves"]):
                if cj not in achados or os.path.getmtime(p) > os.path.getmtime(achados[cj]):
                    achados[cj] = p
    return achados


def _num(v, padrao=0.0):
    if isinstance(v, bool) or v is None or v == "":
        return padrao
    try:
        return float(v)
    except (TypeError, ValueError):
        return padrao


def _txt(v):
    return str(v).strip() if v not in (None, "") else ""


# ───────────────────────── leitura da planilha ─────────────────────────
def _codigo(valor, tipo_fill, rgb, atual):
    if isinstance(valor, (int, float)) and not isinstance(valor, bool):
        return "2" if int(valor) == atual else "1"
    if tipo_fill == "solid" and rgb == COR_VERMELHA:
        return "3"
    if tipo_fill == "lightUp" or tipo_fill is None:
        return "4"
    return "0"


def parse(path):
    import openpyxl
    from openpyxl.utils import column_index_from_string as CI, get_column_letter as L
    wb = openpyxl.load_workbook(path, data_only=True)
    ws = wb["PLS-Habitação"]
    atual = int(_num(ws["F6"].value, 0))

    # unidades: colunas a partir de BD com nome na linha 20; junta (linha 18) e pavimento (linha 19) acompanham
    unidades, junta, pav = [], "", ""
    for c in range(CI("BD"), ws.max_column + 1):
        h18, h19, h20 = ws.cell(18, c).value, ws.cell(19, c).value, ws.cell(20, c).value
        if h18 not in (None, ""):
            junta = _txt(h18)
        if h19 not in (None, ""):
            pav = _txt(h19)
        if h20 in (None, ""):
            continue
        unidades.append({"c": L(c), "j": junta, "p": pav, "u": _txt(h20)})
    idx_col = {u["c"]: i for i, u in enumerate(unidades)}
    c_ini, c_fim = CI(unidades[0]["c"]), CI(unidades[-1]["c"])

    # células mescladas na grade: (linha, coluna) -> coluna âncora
    ancora = {}
    for m in ws.merged_cells.ranges:
        if m.min_col >= c_ini and m.max_col <= c_fim and m.min_row >= 21:
            for r in range(m.min_row, m.max_row + 1):
                for c in range(m.min_col, m.max_col + 1):
                    ancora[(r, c)] = (m.min_row, m.min_col)

    macros = {}
    servicos = []
    for r in range(21, ws.max_row + 1):
        item = _txt(ws.cell(r, 1).value)
        if not item or not re.match(r"^\d+(\.\d+)*$", item):
            continue
        if re.match(r"^\d+$", item):
            macros[int(item)] = _txt(ws.cell(r, 2).value).title()
            continue
        tipo, quant = _txt(ws.cell(r, 3).value), _num(ws.cell(r, 4).value)
        if not tipo or quant <= 0:
            continue
        codigos, merge, tem_grade, por_med = [], {}, False, {}
        for u in unidades:
            c = CI(u["c"])
            ar, ac = ancora.get((r, c), (r, c))
            cel = ws.cell(ar, ac)
            f = cel.fill
            rgb = f.fgColor.rgb if (f.fgColor is not None and f.fgColor.type == "rgb") else None
            cod = _codigo(cel.value, f.fill_type, rgb, atual)
            if cod != "4":
                tem_grade = True
            codigos.append(cod)
            if isinstance(cel.value, (int, float)) and not isinstance(cel.value, bool) and (ar, ac) == (r, c):
                por_med[str(int(cel.value))] = por_med.get(str(int(cel.value)), 0) + 1
            if (ar, ac) != (r, c):
                merge[u["c"]] = L(ac)
        if not tem_grade:
            continue
        macro_item = int(item.split(".")[0])
        servicos.append({
            "id": r, "row": r, "item": item, "nome": _txt(ws.cell(r, 2).value), "tipo": tipo, "quant": quant,
            "macroItem": macro_item, "valor": round(_num(ws.cell(r, 9).value), 2),
            "pct": round(_num(ws.cell(r, 16).value), 6), "medido_rs": round(_num(ws.cell(r, 17).value), 2),
            "s": "".join(codigos), "merge": merge, "por_med": por_med,
        })
    for s in servicos:
        s["macro"] = macros.get(s["macroItem"], "Item %d" % s["macroItem"])

    # resumo global (aba 'Medição Global', linha 'TOTAL')
    glob_ = {}
    if "Medição Global" in wb.sheetnames:
        wg = wb["Medição Global"]
        for r in range(17, 40):
            if _txt(wg.cell(r, 2).value).upper() == "TOTAL":
                glob_ = {
                    "custo_total": round(_num(wg.cell(r, 4).value), 2),
                    "pct_anterior": _num(wg.cell(r, 6).value), "pct_periodo": _num(wg.cell(r, 7).value),
                    "pct_acumulado": _num(wg.cell(r, 8).value),
                    "rs_anterior": round(_num(wg.cell(r, 9).value), 2), "rs_periodo": round(_num(wg.cell(r, 10).value), 2),
                    "rs_acumulado": round(_num(wg.cell(r, 11).value), 2),
                }
                break
    hab = {"orcamento_habitacao": round(_num(ws["H8"].value), 2), "orcamento_global": round(_num(ws["H14"].value), 2)}
    if not glob_.get("custo_total"):
        glob_["custo_total"] = hab["orcamento_global"]
    periodo = ws["F10"].value
    obra = _txt(ws["A12"].value)
    wb.close()
    ultima = max([int(k) for sv in servicos for k in (sv.get("por_med") or {})] or [0])
    return {
        # medicao_atual = última medição JÁ medida pela Caixa na planilha; a seguinte é a que se mede no painel.
        # (F6 traz o nº da medição em preparo: na planilha em branco é a 11 e as células só vão até a 10.)
        "medicao_atual": ultima or atual, "medicao_periodo": atual, "obra": obra,
        "periodo": _txt(periodo), "unidades": unidades, "servicos": servicos,
        "macros": {str(k): v for k, v in macros.items()}, "global": glob_, **hab,
    }


def _data_br(txt):
    m = re.match(r"\s*(\d{2})/(\d{2})/(\d{4})", str(txt or ""))
    return datetime.date(int(m.group(3)), int(m.group(2)), int(m.group(1))) if m else None


def _soma_meses(d, n):
    import calendar
    y, mth = divmod(d.year * 12 + (d.month - 1) + n, 12)
    return datetime.date(y, mth + 1, min(d.day, calendar.monthrange(y, mth + 1)[1]))


def periodos_medicoes(periodo_txt, atual):
    """[{n, ini, fim}] de 1 até a próxima medição — o período da medição atual vem escrito na PLS ('10/08/2026 a 10/09/2026')."""
    partes = re.findall(r"\d{2}/\d{2}/\d{4}", str(periodo_txt or ""))
    if len(partes) < 2 or not atual:
        return []
    fim_atual = _data_br(partes[1])
    out = []
    for n in range(1, atual + 2):
        fim = _soma_meses(fim_atual, n - atual)
        ini = _soma_meses(fim_atual, n - 1 - atual) + datetime.timedelta(days=1)
        out.append({"n": n, "ini": ini.isoformat(), "fim": fim.isoformat()})
    return out


# ───────────────────────── cache (a leitura leva dezenas de segundos) ─────────────────────────
def _cache_path(pasta, cj):
    return os.path.join(pasta, ".pls_cache_%s.json" % cj)


def _ler_cache(pasta, cj, mtime):
    try:
        with open(_cache_path(pasta, cj), encoding="utf-8") as fh:
            c = json.load(fh)
        if abs(c.get("mtime", 0) - mtime) < 1 and c.get("versao") == 5:
            return c["dados"]
    except Exception:
        pass
    return None


def _processar(pasta, cj, path, mtime):
    try:
        dados = parse(path)
        dados["arquivo"] = os.path.basename(path)
        dados["atualizado_em"] = datetime.datetime.fromtimestamp(mtime).strftime("%d/%m/%Y %H:%M")
        tmp = _cache_path(pasta, cj) + ".tmp"
        with open(tmp, "w", encoding="utf-8") as fh:
            json.dump({"mtime": mtime, "versao": 5, "dados": dados}, fh, ensure_ascii=False)
        os.replace(tmp, _cache_path(pasta, cj))
        _cache_mem[cj] = {"mtime": mtime, "dados": dados}
        _parse_threads[cj]["erro"] = ""
    except Exception as exc:  # pragma: no cover
        _parse_threads[cj]["erro"] = "%s: %s" % (type(exc).__name__, exc)


def carregar(pasta, cj):
    """Dados do conjunto, ou {'ok': False, 'carregando': True} enquanto a planilha é lida em segundo plano."""
    achados = localizar(pasta)
    path = achados.get(cj)
    if not path:
        return {"ok": False, "error": "Planilha PLS do %s não encontrada em dados/uau (PLS_HIS_MED*...xlsx)." % cj.upper()}
    mtime = os.path.getmtime(path)
    mem = _cache_mem.get(cj)
    if mem and abs(mem["mtime"] - mtime) < 1:
        return {"ok": True, **mem["dados"]}
    dados = _ler_cache(pasta, cj, mtime)
    if dados:
        _cache_mem[cj] = {"mtime": mtime, "dados": dados}
        return {"ok": True, **dados}
    with _parse_lock:
        th = _parse_threads.get(cj)
        if not th or not th["thread"].is_alive():
            if th and th.get("erro") and th.get("mtime") == mtime:
                return {"ok": False, "error": "Falha ao ler a planilha PLS: " + th["erro"]}
            t = threading.Thread(target=_processar, args=(pasta, cj, path, mtime), daemon=True)
            _parse_threads[cj] = {"thread": t, "erro": "", "mtime": mtime}
            t.start()
    return {"ok": False, "carregando": True, "message": "Lendo a planilha PLS (na primeira vez leva cerca de 1 minuto)…"}


# ───────────────────────── marcações da medição seguinte ─────────────────────────
def _estado_path(pasta):
    return os.path.join(pasta, "pls_medicao_estado.json")


def _estado_ler(pasta):
    try:
        with open(_estado_path(pasta), encoding="utf-8") as fh:
            e = json.load(fh)
        return e if isinstance(e, dict) else {}
    except Exception:
        return {}


def _estado_gravar(pasta, e):
    tmp = _estado_path(pasta) + ".tmp"
    with open(tmp, "w", encoding="utf-8") as fh:
        json.dump(e, fh, ensure_ascii=False, indent=1)
    os.replace(tmp, _estado_path(pasta))


def estado_do_conjunto(pasta, cj, base):
    """Marcações do conjunto; se a planilha trocou de medição, as marcações antigas vão para o arquivo morto."""
    with _lock_estado:
        e = _estado_ler(pasta)
        c = e.get(cj) or {}
        if c.get("base") != base:
            if c.get("marks"):
                c.setdefault("arquivadas", []).append({"base": c.get("base"), "marks": c.get("marks"),
                                                       "em": time.strftime("%Y-%m-%d %H:%M:%S")})
            c = {"base": base, "marks": {}, "log": c.get("log", [])[-200:], "arquivadas": c.get("arquivadas", [])}
            e[cj] = c
            _estado_gravar(pasta, e)
        return c


def _ancoras_marcaveis(servico, unidades):
    """colunas âncora que aceitam marcação (em aberto ou glosadas)."""
    ok = set()
    for i, u in enumerate(unidades):
        col = u["c"]
        if servico["merge"].get(col, col) != col:
            continue
        if servico["s"][i] in ("0", "3"):
            ok.add(col)
    return ok


def marcar(pasta, cj, linha, colunas, ligar, usuario=""):
    """Liga/desliga a medição seguinte de unidades de um serviço. Devolve {'ok', 'marks'}."""
    dados = carregar(pasta, cj)
    if not dados.get("ok"):
        return {"ok": False, "error": dados.get("error") or "Planilha ainda carregando."}
    serv = next((s for s in dados["servicos"] if s["id"] == int(linha)), None)
    if not serv:
        return {"ok": False, "error": "Serviço não encontrado."}
    validas = _ancoras_marcaveis(serv, dados["unidades"])
    alvo = [serv["merge"].get(c, c) for c in (colunas or [])]
    alvo = [c for c in dict.fromkeys(alvo) if c in validas]
    with _lock_estado:
        e = _estado_ler(pasta)
        c = e.get(cj)
        if not c or c.get("base") != dados["medicao_atual"]:
            c = estado_do_conjunto(pasta, cj, dados["medicao_atual"])
            e = _estado_ler(pasta)
            c = e.get(cj, c)
        atuais = set(c["marks"].get(str(linha), []))
        antes = set(atuais)
        if ligar:
            atuais |= set(alvo)
        else:
            atuais -= set(alvo)
        if atuais:
            c["marks"][str(linha)] = sorted(atuais, key=lambda x: (len(x), x))
        else:
            c["marks"].pop(str(linha), None)
        if atuais != antes:
            c.setdefault("log", []).append({"em": time.strftime("%Y-%m-%d %H:%M:%S"), "por": usuario, "linha": int(linha),
                                            "item": serv["item"], "ligar": bool(ligar), "qtd": len(atuais ^ antes)})
            c["log"] = c["log"][-500:]
        e[cj] = c
        _estado_gravar(pasta, e)
        return {"ok": True, "marks": c["marks"]}


def limpar(pasta, cj, linha=None):
    with _lock_estado:
        e = _estado_ler(pasta)
        c = e.get(cj)
        if not c:
            return {"ok": True, "marks": {}}
        if linha is None:
            c["marks"] = {}
        else:
            c["marks"].pop(str(linha), None)
        e[cj] = c
        _estado_gravar(pasta, e)
        return {"ok": True, "marks": c["marks"]}


def _por_evento(serv):
    return serv["valor"] / serv["quant"] if serv["quant"] else 0.0


def resumo(pasta):
    """Resumo por conjunto para o Painel BI e para o cabeçalho da aba."""
    out = {"ok": True, "conjuntos": {}}
    for cj, cfg in CONJUNTOS.items():
        d = carregar(pasta, cj)
        if not d.get("ok"):
            out["conjuntos"][cj] = {"ok": False, "nome": cfg["nome"], "carregando": bool(d.get("carregando")),
                                    "error": d.get("error", "")}
            continue
        c = estado_do_conjunto(pasta, cj, d["medicao_atual"])
        marcado_rs, marcado_ev = 0.0, 0
        por_servico = []
        for s in d["servicos"]:
            cols = c["marks"].get(str(s["id"]), [])
            if cols:
                v = _por_evento(s) * len(cols)
                marcado_rs += v
                marcado_ev += len(cols)
                por_servico.append({"item": s["item"], "nome": s["nome"], "eventos": len(cols), "valor": round(v, 2)})
        g = d.get("global", {})
        total = g.get("custo_total") or d.get("orcamento_global") or 0.0
        acum = g.get("rs_acumulado", 0.0)
        out["conjuntos"][cj] = {
            "ok": True, "nome": cfg["nome"], "medicao_atual": d["medicao_atual"], "proxima": d["medicao_atual"] + 1,
            "periodo": d.get("periodo", ""), "arquivo": d.get("arquivo", ""), "custo_total": total,
            "medido_acumulado": acum, "pct_acumulado": g.get("pct_acumulado", 0.0),
            "medido_periodo": g.get("rs_periodo", 0.0), "pct_periodo": g.get("pct_periodo", 0.0),
            "marcado_valor": round(marcado_rs, 2), "marcado_eventos": marcado_ev,
            "marcado_pct": round(marcado_rs / total, 6) if total else 0.0,
            "projetado_valor": round(acum + marcado_rs, 2),
            "projetado_pct": round((acum + marcado_rs) / total, 6) if total else 0.0,
            "servicos_marcados": sorted(por_servico, key=lambda x: -x["valor"])[:12],
            "ultima_marcacao": (c.get("log") or [{}])[-1].get("em", ""),
        }
    return out


# ───────────────────────── exportação ─────────────────────────
def exportar_xlsx(pasta, cj):
    """XLSX no formato da PLS (mesma linha/coluna da planilha oficial) + relatório das unidades marcadas."""
    import io
    import openpyxl
    from openpyxl.styles import Font
    d = carregar(pasta, cj)
    if not d.get("ok"):
        return None, d.get("error") or "Planilha ainda carregando."
    c = estado_do_conjunto(pasta, cj, d["medicao_atual"])
    prox = d["medicao_atual"] + 1
    rotulo = {u["c"]: u for u in d["unidades"]}
    wb = openpyxl.Workbook()
    simu = wb.active
    simu.title = "PLS_SIMU"
    rel = wb.create_sheet("Relatório")
    rel.append(["Item", "Serviço", "Junta", "Pavimento", "Unidade", "Coluna PLS", "Linha PLS", "Valor do evento (R$)"])
    for cel in rel[1]:
        cel.font = Font(bold=True)
    total, n = 0.0, 0
    for s in d["servicos"]:
        cols = c["marks"].get(str(s["id"]), [])
        if not cols:
            continue
        simu["A%d" % s["row"]] = s["item"]
        simu["B%d" % s["row"]] = s["nome"]
        for col in cols:
            simu["%s%d" % (col, s["row"])] = prox
            u = rotulo.get(col, {})
            v = round(_por_evento(s), 2)
            rel.append([s["item"], s["nome"], u.get("j", ""), u.get("p", ""), u.get("u", ""), col, s["row"], v])
            total += v
            n += 1
    if not n:
        return None, "Nenhuma unidade marcada para a medição %d." % prox
    rel.append([])
    rel.append(["", "TOTAL", "", "", "", "", "", round(total, 2)])
    for col, w in zip("ABCDEFGH", (10, 48, 14, 12, 26, 11, 10, 20)):
        rel.column_dimensions[col].width = w
    bio = io.BytesIO()
    wb.save(bio)
    bio.seek(0)
    return bio, "PLS_SIMU-medicao-%d-%s.xlsx" % (prox, cj)


# ───────────────────────── relatório (PDF / Excel) ─────────────────────────
def _brl(v):
    return ("R$ {:,.2f}".format(float(v or 0))).replace(",", "X").replace(".", ",").replace("X", ".")


def _pctf(v):
    return ("{:.2f}%".format(float(v or 0) * 100)).replace(".", ",")


def _montar_relatorio(pasta, cj):
    """Todos os números do relatório de um conjunto (a mesma base do painel)."""
    d = carregar(pasta, cj)
    if not d.get("ok"):
        return None
    c = estado_do_conjunto(pasta, cj, d["medicao_atual"])
    unid = d["unidades"]
    servicos, marcadas = [], []
    for s in d["servicos"]:
        cont = {"0": 0, "1": 0, "2": 0, "3": 0}
        ancoras = []
        for i, u in enumerate(unid):
            if s["merge"].get(u["c"], u["c"]) != u["c"]:
                continue
            cod = s["s"][i]
            if cod in cont:
                cont[cod] += 1
            ancoras.append((u, cod))
        marc = set(c["marks"].get(str(s["id"]), []))
        validas = [(u, cod) for u, cod in ancoras if u["c"] in marc and cod in ("0", "3")]
        total = sum(cont.values())
        medido = cont["1"] + cont["2"]
        por_ev = _por_evento(s)
        servicos.append({
            "item": s["item"], "nome": s["nome"], "macro": s["macro"], "macroItem": s["macroItem"], "tipo": s["tipo"],
            "eventos": total, "medido": medido, "aberto": cont["0"], "glosado": cont["3"], "agora": len(validas),
            "valor_servico": s["valor"], "valor_agora": round(por_ev * len(validas), 2),
            "pct_apos": (medido + len(validas)) / total if total else 0.0,
        })
        for u, cod in validas:
            marcadas.append({"item": s["item"], "nome": s["nome"], "macro": s["macro"], "junta": u["j"], "pav": u["p"],
                             "unidade": u["u"], "coluna": u["c"], "linha": s["row"], "glosa": cod == "3",
                             "valor": round(por_ev, 2)})
    r = resumo(pasta)["conjuntos"][cj]
    return {"cj": cj, "dados": d, "resumo": r, "servicos": servicos, "marcadas": marcadas}


def _agrupar_locais(itens):
    """[('JUNTA A', 'PAV 2', ['210 -B', '209 -B']), ...] preservando a ordem."""
    out = {}
    for m in itens:
        out.setdefault((m["junta"], m["pav"]), []).append(m["unidade"] + (" (glosa)" if m["glosa"] else ""))
    return [(j, p, us) for (j, p), us in out.items()]


def relatorio_xlsx(pasta, conjuntos, usuario=""):
    import io
    import openpyxl
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter
    blocos = [b for b in (_montar_relatorio(pasta, cj) for cj in conjuntos) if b]
    if not blocos:
        return None, "Planilhas PLS ainda carregando ou não encontradas."
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    neg = Font(bold=True)
    cab = PatternFill("solid", fgColor="A50E24")
    cab_font = Font(bold=True, color="FFFFFF")
    moeda = "#,##0.00"

    def cabecalho(ws, titulos, larguras):
        ws.append(titulos)
        for cel in ws[ws.max_row]:
            cel.font, cel.fill = cab_font, cab
            cel.alignment = Alignment(wrap_text=True, vertical="center")
        for i, w in enumerate(larguras):
            ws.column_dimensions[get_column_letter(i + 1)].width = w
        ws.freeze_panes = "A%d" % (ws.max_row + 1)

    if len(blocos) > 1:
        ws = wb.create_sheet("Consolidado")
        ws["A1"] = "Medição PLS — consolidado"
        ws["A1"].font = Font(bold=True, size=14)
        ws["A2"] = "Gerado em %s por %s" % (time.strftime("%d/%m/%Y %H:%M"), usuario or "-")
        ws.append([])
        cabecalho(ws, ["Conjunto", "Custo total", "Medido acumulado", "% acumulado", "Medido agora", "Projetado", "% projetado"],
                  [30, 18, 18, 14, 18, 18, 14])
        tot = [0.0, 0.0, 0.0, 0.0]
        for b in blocos:
            r = b["resumo"]
            ws.append([r["nome"], r["custo_total"], r["medido_acumulado"], r["pct_acumulado"], r["marcado_valor"],
                       r["projetado_valor"], r["projetado_pct"]])
            tot = [tot[0] + r["custo_total"], tot[1] + r["medido_acumulado"], tot[2] + r["marcado_valor"], tot[3] + r["projetado_valor"]]
        ws.append(["TOTAL", tot[0], tot[1], tot[1] / tot[0] if tot[0] else 0, tot[2], tot[3], tot[3] / tot[0] if tot[0] else 0])
        for cel in ws[ws.max_row]:
            cel.font = neg
        for row in ws.iter_rows(min_row=5):
            for k in (1, 2, 4, 5):
                row[k].number_format = moeda
            for k in (3, 6):
                row[k].number_format = "0.00%"
    for b in blocos:
        r, d, cj = b["resumo"], b["dados"], b["cj"].upper()
        ws = wb.create_sheet("%s Resumo" % cj)
        ws["A1"] = "Medição PLS — %s" % r["nome"]
        ws["A1"].font = Font(bold=True, size=14)
        linhas = [("Empreendimento", d.get("obra", "")), ("Período de referência", d.get("periodo", "")),
                  ("Medição atual (planilha)", d["medicao_atual"]), ("Medindo agora", "Medição %d" % r["proxima"]),
                  ("Planilha de origem", d.get("arquivo", "")), ("Gerado em", time.strftime("%d/%m/%Y %H:%M")),
                  ("Gerado por", usuario or "-"), ("", ""), ("Custo total do contrato", r["custo_total"]),
                  ("Medido acumulado (Caixa)", r["medido_acumulado"]), ("% acumulado (Caixa)", r["pct_acumulado"]),
                  ("Medido agora (marcado no painel)", r["marcado_valor"]), ("Unidades marcadas", r["marcado_eventos"]),
                  ("Acumulado projetado", r["projetado_valor"]), ("% projetado", r["projetado_pct"])]
        for i, (k, v) in enumerate(linhas, start=3):
            ws.cell(i, 1, k).font = neg
            cel = ws.cell(i, 2, v)
            if k.startswith("%"):
                cel.number_format = "0.00%"
            elif isinstance(v, float):
                cel.number_format = moeda
        ws.column_dimensions["A"].width = 36
        ws.column_dimensions["B"].width = 46

        ws = wb.create_sheet("%s Serviços" % cj)
        cabecalho(ws, ["Item", "Serviço", "Grupo", "Tipo", "Eventos", "Medido (Caixa)", "Em aberto", "Glosado", "Medido agora",
                       "% após medição", "Valor do serviço (R$)", "Valor medido agora (R$)"],
                  [10, 52, 28, 12, 9, 13, 11, 10, 13, 13, 20, 20])
        for s in b["servicos"]:
            ws.append([s["item"], s["nome"], s["macro"], s["tipo"], s["eventos"], s["medido"], s["aberto"], s["glosado"], s["agora"],
                       s["pct_apos"], s["valor_servico"], s["valor_agora"]])
            ws.cell(ws.max_row, 10).number_format = "0.00%"
            ws.cell(ws.max_row, 11).number_format = moeda
            ws.cell(ws.max_row, 12).number_format = moeda
            if s["agora"]:
                for cel in ws[ws.max_row]:
                    cel.fill = PatternFill("solid", fgColor="FFF4C2")

        ws = wb.create_sheet("%s Marcadas" % cj)
        cabecalho(ws, ["Item", "Serviço", "Junta", "Pavimento", "Unidade", "Coluna PLS", "Linha PLS", "Glosa remedida",
                       "Valor do evento (R$)"], [10, 52, 16, 12, 28, 11, 10, 14, 20])
        for m in b["marcadas"]:
            ws.append([m["item"], m["nome"], m["junta"], m["pav"], m["unidade"], m["coluna"], m["linha"],
                       "sim" if m["glosa"] else "", m["valor"]])
            ws.cell(ws.max_row, 9).number_format = moeda
        ws.append([])
        ws.append(["", "TOTAL", "", "", "", "", "", "", sum(m["valor"] for m in b["marcadas"])])
        ws.cell(ws.max_row, 9).number_format = moeda
        for cel in ws[ws.max_row]:
            cel.font = neg
    bio = io.BytesIO()
    wb.save(bio)
    bio.seek(0)
    nome = "Relatorio_PLS_medicao-%d_%s.xlsx" % (blocos[0]["resumo"]["proxima"], "_".join(b["cj"] for b in blocos))
    return bio, nome


def relatorio_pdf(pasta, conjuntos, usuario=""):
    import io
    from fpdf import FPDF
    blocos = [b for b in (_montar_relatorio(pasta, cj) for cj in conjuntos) if b]
    if not blocos:
        return None, "Planilhas PLS ainda carregando ou não encontradas."

    def seg(t):  # as fontes padrão do PDF são latin-1 (o português cabe); troca o que não couber
        t = str(t if t not in (None, "") else "-")
        for a, b_ in (("—", "-"), ("–", "-"), ("•", "-")):
            t = t.replace(a, b_)
        return t.encode("latin-1", "replace").decode("latin-1")

    class PDF(FPDF):
        def footer(self):
            self.set_y(-10)
            self.set_font("Helvetica", "", 7)
            self.set_text_color(120, 120, 120)
            self.cell(0, 5, seg("Medicao PLS - gerado em %s - pagina %d" % (time.strftime("%d/%m/%Y %H:%M"), self.page_no())), align="C")

    pdf = PDF(orientation="L", unit="mm", format="A4")
    pdf.set_auto_page_break(auto=True, margin=14)
    VERM = (165, 14, 36)

    def titulo(txt, tam=14):
        pdf.set_font("Helvetica", "B", tam)
        pdf.set_text_color(*VERM)
        pdf.cell(0, 8, seg(txt), new_x="LMARGIN", new_y="NEXT")
        pdf.set_text_color(30, 30, 30)

    def linha_tabela(cols, larguras, negrito=False, fundo=None, alinhar=None, h=5.2, tam=7.5):
        pdf.set_font("Helvetica", "B" if negrito else "", tam)
        if fundo:
            pdf.set_fill_color(*fundo)
        for i, (t, w) in enumerate(zip(cols, larguras)):
            pdf.cell(w, h, seg(t)[: int(w * 2.1)], border=1, fill=bool(fundo), align=(alinhar[i] if alinhar else "L"))
        pdf.ln(h)

    for b in blocos:
        r, d = b["resumo"], b["dados"]
        pdf.add_page()
        titulo("Medicao PLS - %s" % r["nome"], 16)
        pdf.set_font("Helvetica", "", 9)
        pdf.set_text_color(90, 90, 90)
        pdf.cell(0, 5, seg("%s - periodo %s - medicao atual %s (planilha) - medindo a medicao %d" %
                           (d.get("obra", ""), d.get("periodo", ""), d["medicao_atual"], r["proxima"])), new_x="LMARGIN", new_y="NEXT")
        pdf.cell(0, 5, seg("Planilha: %s - gerado por %s" % (d.get("arquivo", ""), usuario or "-")), new_x="LMARGIN", new_y="NEXT")
        pdf.set_text_color(30, 30, 30)
        pdf.ln(3)
        larg = [100, 45, 30]
        linha_tabela(["Resumo financeiro", "Valor (R$)", "% contrato"], larg, True, (240, 232, 232), ["L", "R", "R"], 6, 8.5)
        for rot, v, p in (("Custo total do contrato", r["custo_total"], None),
                          ("Medido acumulado pela Caixa (ate a medicao %d)" % d["medicao_atual"], r["medido_acumulado"], r["pct_acumulado"]),
                          ("Medido agora (medicao %d) - %d unidades" % (r["proxima"], r["marcado_eventos"]), r["marcado_valor"], r["marcado_pct"]),
                          ("Acumulado projetado", r["projetado_valor"], r["projetado_pct"])):
            linha_tabela([rot, _brl(v), _pctf(p) if p is not None else ""], larg, rot.startswith("Acumulado"), None, ["L", "R", "R"], 6, 8.5)
        pdf.ln(4)

        marc = [s for s in b["servicos"] if s["agora"]]
        titulo("Medido agora - medicao %d (%d servicos)" % (r["proxima"], len(marc)), 12)
        if not marc:
            pdf.set_font("Helvetica", "I", 9)
            pdf.cell(0, 6, seg("Nenhuma unidade marcada ate o momento."), new_x="LMARGIN", new_y="NEXT")
        locais = {}
        for m in b["marcadas"]:
            locais.setdefault(m["item"] + "|" + m["nome"], []).append(m)
        for s in marc:
            itens = locais.get(s["item"] + "|" + s["nome"], [])
            linha_tabela([s["item"], s["nome"], "%d un." % s["agora"], _brl(s["valor_agora"]), _pctf(s["pct_apos"])],
                         [16, 150, 20, 32, 20], True, (245, 238, 214), ["L", "L", "R", "R", "R"])
            for j, p, us in _agrupar_locais(itens):
                pdf.set_font("Helvetica", "", 7)
                pdf.set_text_color(70, 70, 70)
                pdf.set_x(pdf.l_margin + 6)
                pdf.multi_cell(0, 4, seg("%s - %s: %s" % (j, p, ", ".join(us))), new_x="LMARGIN", new_y="NEXT")
                pdf.set_text_color(30, 30, 30)
        pdf.add_page()
        titulo("Situacao geral dos servicos - %s" % r["nome"], 12)
        larg2 = [16, 108, 22, 18, 18, 18, 18, 18, 22, 20]
        cab = ["Item", "Servico", "Tipo", "Eventos", "Medido", "Aberto", "Glosado", "Agora", "% apos", "Valor agora"]
        al = ["L", "L", "L", "R", "R", "R", "R", "R", "R", "R"]
        linha_tabela(cab, larg2, True, (240, 232, 232), al)
        macro_atual = None
        for s in b["servicos"]:
            if pdf.get_y() > 190:
                pdf.add_page()
                linha_tabela(cab, larg2, True, (240, 232, 232), al)
            if s["macroItem"] != macro_atual:
                macro_atual = s["macroItem"]
                linha_tabela(["%d" % s["macroItem"], s["macro"].upper()] + [""] * 8, larg2, True, (232, 232, 240), al)
            linha_tabela([s["item"], s["nome"], s["tipo"], s["eventos"], s["medido"], s["aberto"], s["glosado"], s["agora"],
                          _pctf(s["pct_apos"]), _brl(s["valor_agora"]) if s["agora"] else ""], larg2, False,
                         (255, 247, 220) if s["agora"] else None, al)
    bio = io.BytesIO(bytes(pdf.output()))
    nome = "Relatorio_PLS_medicao-%d_%s.pdf" % (blocos[0]["resumo"]["proxima"], "_".join(b["cj"] for b in blocos))
    return bio, nome


# ───────────────────────── de-para PLS ↔ orçamento (EAP) e físico-financeiro ─────────────────────────
# itens do orçamento (nível 3 da EAP do UAU, igual nos dois conjuntos) que recebem medição da PLS
ORC_ITENS = [
    ("02.01", "SERVICOS PRELIMINARES DE OBRA"), ("02.02", "FUNDAÇÃO/ CONTENÇÃO"), ("02.03", "ESTRUTURAS"),
    ("02.04", "ALVENARIAS DE VEDAÇÃO"), ("02.05", "ESQUADRIAS E GUARDACORPOS"), ("02.06", "INSTALAÇÕES"),
    ("02.07", "LOUÇAS METAIS BANCADAS E ACESSORIOS"), ("02.08", "IMPERMEABILIZAÇÕES"),
    ("02.09", "REVESTIMENTOS INTERNOS BASICOS"), ("02.10", "REVESTIMENTOS EXTERNOS"),
    ("02.11", "ELEMENTOS ESPECIAIS E COMPLEMENTOS"), ("02.12", "REVESTIMENTOS ACABAMENTOS"), ("02.13", "FORROS"),
    ("02.14", "COBERTURA"), ("02.15", "PAVIMENTAÇÕES"), ("02.16", "PISOS"), ("02.17", "PINTURAS INTERNAS"),
    ("02.18", "PINTURAS EXTERNAS"), ("02.19", "PAISAGISMO"), ("02.20", "LIGACOES DEFINITIVAS"),
    ("02.21", "FINALIZAÇÃO DE OBRA"), ("02.22", "DESPESAS PÓS ENTREGA"),
]
# conjunto da PLS -> centro de custo do UAU que recebe a medição
CC_DO_CONJUNTO = {"cj03": "116-11602", "cj05": "116-11603"}
# sugestão automática: prefixo do item da PLS -> item do orçamento (vence o prefixo mais específico)
_REGRAS = [
    ("2.1", "02.02"), ("3", "02.03"),
    ("4.1", "02.04"), ("4.2", "02.05"), ("4.3", "02.05"),
    ("5.1", "02.14"), ("5.2", "02.08"),
    ("6.1", "02.09"), ("6.2", "02.12"), ("6.4", "02.13"),
    ("6.5", "02.17"), ("6.5.9", "02.18"), ("6.5.10", "02.18"),
    ("7.2", "02.16"), ("7.2.1", "02.15"), ("7.2.2", "02.15"), ("7.2.3", "02.15"),
    ("7.3", "02.11"), ("7.3.1", "02.16"), ("7.3.2", "02.16"),
    ("7.4", "02.15"),
    ("8", "02.06"), ("8.4", "02.07"),
    ("9", "02.21"),
]


def depara_padrao(item):
    melhor, tam = "", -1
    for pref, cod in _REGRAS:
        if (item == pref or item.startswith(pref + ".")) and len(pref) > tam:
            melhor, tam = cod, len(pref)
    return melhor


def _depara_path(pasta):
    return os.path.join(pasta, "pls_depara.json")


def depara_ler(pasta):
    try:
        with open(_depara_path(pasta), encoding="utf-8") as fh:
            d = json.load(fh)
        return d if isinstance(d, dict) else {}
    except Exception:
        return {}


def depara_efetivo(pasta, itens):
    """{item da PLS: código do orçamento} = sugestão automática + ajustes manuais."""
    manual = depara_ler(pasta).get("manual", {})
    return {i: (manual[i] if i in manual else depara_padrao(i)) for i in itens}


def depara_tabela(pasta):
    """Linhas do de-para (usa o CJ03; a estrutura de serviços é a mesma nos dois conjuntos)."""
    d = carregar(pasta, "cj03")
    if not d.get("ok"):
        d = carregar(pasta, "cj05")
    if not d.get("ok"):
        return d
    manual = depara_ler(pasta).get("manual", {})
    vistos, linhas = set(), []
    for s in d["servicos"]:
        if s["item"] in vistos:
            continue
        vistos.add(s["item"])
        linhas.append({"item": s["item"], "nome": s["nome"], "macroItem": s["macroItem"], "macro": s["macro"],
                       "valor": s["valor"], "orc": manual[s["item"]] if s["item"] in manual else depara_padrao(s["item"]),
                       "manual": s["item"] in manual, "sugerido": depara_padrao(s["item"])})
    return {"ok": True, "linhas": linhas, "orc_itens": [{"cod": c, "nm": n} for c, n in ORC_ITENS]}


def depara_salvar(pasta, item, cod):
    """cod = '' (sem correspondência), código do orçamento, ou None para voltar à sugestão automática."""
    cods = {c for c, _ in ORC_ITENS}
    if cod not in (None, "") and cod not in cods:
        return {"ok": False, "error": "Item de orçamento inválido."}
    with _lock_estado:
        d = depara_ler(pasta)
        man = d.setdefault("manual", {})
        if cod is None or (cod == depara_padrao(item)):
            man.pop(item, None)
        else:
            man[item] = cod
        tmp = _depara_path(pasta) + ".tmp"
        with open(tmp, "w", encoding="utf-8") as fh:
            json.dump(d, fh, ensure_ascii=False, indent=1)
        os.replace(tmp, _depara_path(pasta))
    return {"ok": True}


def depara_restaurar(pasta):
    with _lock_estado:
        d = depara_ler(pasta)
        d["manual"] = {}
        tmp = _depara_path(pasta) + ".tmp"
        with open(tmp, "w", encoding="utf-8") as fh:
            json.dump(d, fh, ensure_ascii=False, indent=1)
        os.replace(tmp, _depara_path(pasta))
    return {"ok": True}


def _contratar_path(pasta):
    return os.path.join(pasta, "pls_a_contratar.json")


def contratar_ler(pasta):
    """Anotações manuais por item do quadro: {chave: {valor, obs, por, em}} (chave = conjunto|cc|item)."""
    try:
        with open(_contratar_path(pasta), encoding="utf-8") as fh:
            d = json.load(fh)
        return d if isinstance(d, dict) else {}
    except Exception:
        return {}


def contratar_salvar(pasta, chave, valor, obs, usuario=""):
    chave = str(chave or "").strip()
    if not chave or chave.count("|") != 2:
        return {"ok": False, "error": "Item inválido."}
    try:
        valor = round(float(valor or 0), 2)
    except (TypeError, ValueError):
        return {"ok": False, "error": "Valor inválido."}
    obs = str(obs or "").strip()[:600]
    with _lock_estado:
        d = contratar_ler(pasta)
        if valor == 0 and not obs:
            d.pop(chave, None)
        else:
            d[chave] = {"valor": valor, "obs": obs, "por": usuario, "em": time.strftime("%d/%m/%Y %H:%M")}
        tmp = _contratar_path(pasta) + ".tmp"
        with open(tmp, "w", encoding="utf-8") as fh:
            json.dump(d, fh, ensure_ascii=False, indent=1)
        os.replace(tmp, _contratar_path(pasta))
    return {"ok": True}


def fisico_financeiro(pasta):
    """Transforma a medição por eventos da PLS em % físico por item do orçamento.

    Para cada serviço da PLS: % da Caixa = coluna 'TOTAL %' da planilha (eventos medidos ÷ eventos do serviço);
    % projetado = esse % + as unidades marcadas na medição seguinte ÷ eventos do serviço.
    Para cada item do orçamento: média dos serviços ligados a ele, ponderada pelo valor do serviço na PLS.
    O valor medido em R$ é calculado no painel (orçado do item × %), porque o orçado vem do SQL do UAU.
    """
    out = {"ok": True, "conjuntos": {}}
    for cj, cc in CC_DO_CONJUNTO.items():
        d = carregar(pasta, cj)
        if not d.get("ok"):
            out["conjuntos"][cj] = {"ok": False, "carregando": bool(d.get("carregando")), "error": d.get("error", "")}
            continue
        c = estado_do_conjunto(pasta, cj, d["medicao_atual"])
        mapa = depara_efetivo(pasta, {s["item"] for s in d["servicos"]})
        itens = {}
        for s in d["servicos"]:
            cod = mapa.get(s["item"], "")
            if not cod:
                continue
            marc = set(c["marks"].get(str(s["id"]), []))
            validas = _ancoras_marcaveis(s, d["unidades"]) & marc
            quant = s["quant"] or 1
            pc = min(1.0, s["pct"])
            pp = min(1.0, pc + len(validas) / quant)
            it = itens.setdefault(cod, {"valor_pls": 0.0, "caixa": 0.0, "proj": 0.0, "servicos": 0, "marcadas": 0, "detalhe": [], "med": {}})
            for n_med, qtd in (s.get("por_med") or {}).items():
                it["med"][n_med] = it["med"].get(n_med, 0.0) + s["valor"] * qtd / quant
            if validas:
                prox_key = str(d["medicao_atual"] + 1)
                it["med"][prox_key] = it["med"].get(prox_key, 0.0) + s["valor"] * len(validas) / quant
            it["valor_pls"] += s["valor"]
            it["caixa"] += s["valor"] * pc
            it["proj"] += s["valor"] * pp
            it["servicos"] += 1
            it["marcadas"] += len(validas)
            it["detalhe"].append({"item": s["item"], "nome": s["nome"], "valor": s["valor"], "pct_caixa": round(pc, 4),
                                  "pct_proj": round(pp, 4)})
        for cod, it in itens.items():
            v = it["valor_pls"] or 1.0
            it["pct_caixa"] = round(it.pop("caixa") / v, 6)
            it["pct_proj"] = round(it.pop("proj") / v, 6)
            it["med"] = {k: round(v_ / (it["valor_pls"] or 1.0), 6) for k, v_ in it["med"].items()}
            it["valor_pls"] = round(it["valor_pls"], 2)
        out["conjuntos"][cj] = {"ok": True, "cc": cc, "nome": CONJUNTOS[cj]["nome"], "medicao_atual": d["medicao_atual"],
                                "proxima": d["medicao_atual"] + 1, "itens": itens,
                                "periodos": periodos_medicoes(d.get("periodo"), d.get("medicao_periodo") or d["medicao_atual"])}
    return out


# ───────────────────────── planilha de auditoria (custo lançado, contratos, OCs, adiantamento) ─────────────────────────
def auditoria_xlsx(payload, usuario=""):
    """Monta o Excel de auditoria a partir do que o painel calculou (JSON enviado pela tela).
    Em cada conjunto: Resumo (com fórmulas que somam os títulos e comparam com o painel), Contratos, OCs e Títulos."""
    import io
    import openpyxl
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter as L

    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    cab_fill = PatternFill("solid", fgColor="A50E24")
    cab_font = Font(bold=True, color="FFFFFF")
    sub_fill = PatternFill("solid", fgColor="F3E8E8")
    neg = Font(bold=True)
    moeda = '#,##0.00;[Red]-#,##0.00'
    pct = "0.00%"

    def cabecalho(ws, linha, titulos, larguras=None):
        for i, t in enumerate(titulos, start=1):
            c = ws.cell(linha, i, t)
            c.font, c.fill = cab_font, cab_fill
            c.alignment = Alignment(wrap_text=True, vertical="center", horizontal="center")
        ws.row_dimensions[linha].height = 32
        if larguras:
            for i, w in enumerate(larguras, start=1):
                ws.column_dimensions[L(i)].width = w
        ws.freeze_panes = ws.cell(linha + 1, 4)
        ws.auto_filter.ref = "A%d:%s%d" % (linha, L(len(titulos)), linha)

    per = payload.get("periodo") or {}
    rot_per = ("%s a %s" % (per.get("ini") or "início", per.get("fim") or "hoje")) if (per.get("ini") or per.get("fim")) else "acumulado (sem período)"

    ws = wb.create_sheet("Leia-me")
    linhas = [
        ("Auditoria — Medição PLS × Orçamento (UAU)", None),
        ("Gerado em", time.strftime("%d/%m/%Y %H:%M")), ("Gerado por", usuario or "-"), ("Período considerado", rot_per),
        ("Inclui as unidades marcadas na medição seguinte", "sim" if payload.get("usaMarcas") else "não"), (None, None),
        ("Custo lançado", "Soma dos títulos do UAU (pagos + a pagar) apropriados ao item da EAP. Aba 'Títulos' traz cada título; a coluna 'Custo considerado' é o que entra na conta."),
        ("Grupo do título", "Cada título cai em UM só grupo: 'Contrato' (tem número de contrato) ou 'OC' (compra). Não há soma em dobro."),
        ("Adiantamento", "Parte do custo lançado classificada como adiantada: contratos de fornecimento (todo o lançado), contratos de serviço (só o excesso sobre o avanço físico) e ordens de compra (todo o lançado). NÃO é um valor a mais."),
        ("Estoque financeiro", "Adiantamento ainda não aplicado = (1 − % físico) do adiantamento, limitado ao custo lançado. No período: variação do estoque."),
        ("Valor agregado (VA)", "Orçado × % físico medido pela PLS (via de-para de EAP). Indiretos medem no mesmo % do direto e são rateados 50% para cada conjunto."),
        ("Custo ajustado", "Custo lançado − estoque financeiro."),
        ("Resultado", "VA − custo ajustado."),
        ("Lançado (contratos + OC)", "Contratos pelo valor total (medido ou não) + OCs pelo valor total (rateadas entre as apropriações com títulos) + títulos avulsos. Resultado = orçado − lançado. Colunas V a Z do Resumo; abas 'Contratos' e 'OC' da lógica."),
        ("Conferência", "Em 'Resumo', a coluna 'Σ títulos (fórmula)' soma a aba Títulos e a 'Diferença' compara com o painel (deve ser 0,00)."),
    ]
    for i, (a, b) in enumerate(linhas, start=1):
        if a:
            ws.cell(i, 1, a).font = neg
        if b is not None:
            ws.cell(i, 2, b).alignment = Alignment(wrap_text=True, vertical="top")
    ws["A1"].font = Font(bold=True, size=14)
    ws.column_dimensions["A"].width = 44
    ws.column_dimensions["B"].width = 120

    for cj in payload.get("conjuntos", []):
        nome = str(cj.get("cj", "")).upper()
        t_nome = "%s Títulos" % nome
        # ---- Títulos (base da auditoria)
        wt = wb.create_sheet(t_nome)
        cabecalho(wt, 1, ["Origem (CC)", "Item", "Descrição do item", "Título", "Credor", "Documento", "Grupo", "Contrato", "OC",
                          "Descrição", "Un", "Qtd", "Status", "Data (baixa/venc.)", "Valor apropriado", "Pago", "A pagar",
                          "Fator de rateio", "Custo considerado", "Considerado no período?", "Linha repetida?"],
                  [12, 9, 34, 12, 34, 14, 10, 12, 10, 44, 6, 8, 10, 14, 15, 15, 15, 9, 16, 12, 10])
        for t in cj.get("titulos", []):
            wt.append([t.get("cc"), t.get("item"), t.get("nmItem"), t.get("numtit"), t.get("credor"), t.get("documento"), t.get("grupo"),
                       t.get("contrato"), t.get("oc"), t.get("descr"), t.get("un"), t.get("qtd"), t.get("status"), t.get("data"),
                       t.get("vlr_ap"), t.get("vlr_pg"), t.get("vlr_a_pagar"), t.get("fator", 1), t.get("considerado"),
                       "SIM" if t.get("noPeriodo", True) else "NÃO", "SIM" if t.get("repetida") else ""])
        n_t = wt.max_row
        for r in range(2, n_t + 1):
            for c in (15, 16, 17, 19):
                wt.cell(r, c).number_format = moeda
        wt.cell(n_t + 2, 14, "TOTAL").font = neg
        wt.cell(n_t + 2, 19, "=SUMIFS(S2:S%d,T2:T%d,\"SIM\")" % (max(n_t, 2), max(n_t, 2))).number_format = moeda
        wt.cell(n_t + 2, 19).font = neg

        def tit(col):
            return "'%s'!$%s$2:$%s$%d" % (t_nome, col, col, max(n_t, 2))

        # ---- Resumo por item (com fórmulas)
        wr = wb.create_sheet("%s Resumo" % nome)
        r_cj = cj.get("resumo", {})
        wr.cell(1, 1, "%s — conferência por item" % cj.get("nome", nome)).font = Font(bold=True, size=13)
        wr.cell(2, 1, "Período: %s  |  VA = orçado × %% físico  |  custo ajustado = custo lançado − estoque" % rot_per)
        cabecalho(wr, 4, ["Tipo", "Origem (CC)", "Item", "Descrição", "Orçado", "% físico", "Valor agregado", "Custo lançado (painel)",
                          "Σ títulos (fórmula)", "Diferença", "Contratos (fórmula)", "OCs (fórmula)", "Adiantamento — contratos",
                          "Adiantamento — OC", "Adiantamento total", "Estoque financeiro", "Custo ajustado", "Resultado",
                          "Estimativa ao término", "Sobra (+) / estouro (−)", "Situação",
                          "Contratos (valor total)", "OCs", "Avulsos", "Lançado = contratos + OC + avulsos (fórmula)", "Orçado − lançado (fórmula)"],
                  [10, 12, 9, 38, 15, 9, 15, 17, 17, 12, 16, 16, 16, 16, 16, 16, 16, 16, 17, 16, 24, 17, 16, 16, 19, 19])
        wr.freeze_panes = "E5"
        for i, it in enumerate(cj.get("itens", []), start=5):
            wr.cell(i, 1, {"direto": "Direto", "indireto": "Indireto (rateio 50%)", "sem_pls": "Sem serviço na PLS"}.get(it.get("tipo"), it.get("tipo")))
            wr.cell(i, 2, it.get("cc"))
            wr.cell(i, 3, it.get("cod"))
            wr.cell(i, 4, it.get("nm"))
            wr.cell(i, 5, it.get("orcado"))
            wr.cell(i, 6, it.get("pct") if it.get("tipo") != "sem_pls" else None)
            wr.cell(i, 7, it.get("medido") if it.get("tipo") != "sem_pls" else None)
            wr.cell(i, 8, it.get("custo"))
            wr.cell(i, 9, "=SUMIFS(%s,%s,B%d,%s,C%d,%s,\"SIM\")" % (tit("S"), tit("A"), i, tit("B"), i, tit("T")))
            wr.cell(i, 10, "=H%d-I%d" % (i, i))
            wr.cell(i, 11, "=SUMIFS(%s,%s,B%d,%s,C%d,%s,\"Contrato\",%s,\"SIM\")" % (tit("S"), tit("A"), i, tit("B"), i, tit("G"), tit("T")))
            wr.cell(i, 12, "=SUMIFS(%s,%s,B%d,%s,C%d,%s,\"OC\",%s,\"SIM\")" % (tit("S"), tit("A"), i, tit("B"), i, tit("G"), tit("T")))
            wr.cell(i, 13, it.get("adCt"))
            wr.cell(i, 14, it.get("adOC"))
            wr.cell(i, 15, "=M%d+N%d" % (i, i) if it.get("adBruto") is not None else None)
            wr.cell(i, 16, it.get("naoMedido"))
            wr.cell(i, 17, "=H%d-P%d" % (i, i))
            wr.cell(i, 18, it.get("resultado"))
            wr.cell(i, 19, it.get("previsao"))
            wr.cell(i, 20, "=E%d-S%d" % (i, i))
            wr.cell(i, 21, it.get("situacao"))
            wr.cell(i, 22, it.get("ctTotal"))
            wr.cell(i, 23, it.get("ocTotal"))
            wr.cell(i, 24, it.get("avulsos"))
            wr.cell(i, 25, "=V%d+W%d+X%d" % (i, i, i))
            wr.cell(i, 26, "=E%d-Y%d" % (i, i))
            for c in (5, 7, 8, 9, 10, 11, 12, 13, 14, 15, 16, 17, 18, 19, 20, 22, 23, 24, 25, 26):
                wr.cell(i, c).number_format = moeda
            wr.cell(i, 6).number_format = pct
        fim = 4 + len(cj.get("itens", []))
        tr = fim + 2
        wr.cell(tr, 4, "TOTAL").font = neg
        for c in (5, 7, 8, 9, 10, 11, 12, 13, 14, 15, 16, 17, 18, 19, 20, 22, 23, 24, 25, 26):
            col = L(c)
            cel = wr.cell(tr, c, "=SUM(%s5:%s%d)" % (col, col, max(fim, 5)))
            cel.number_format, cel.font, cel.fill = moeda, neg, sub_fill
        extra = cj.get("semGastoOrc") or 0
        if extra:
            wr.cell(tr + 1, 4, "+ orçados em itens sem gasto e sem medição (estimados ao orçado)")
            wr.cell(tr + 1, 5, extra).number_format = moeda
            wr.cell(tr + 1, 19, cj.get("semGastoEac") or extra).number_format = moeda
        wr.cell(tr + 3, 4, "KPIs do painel").font = neg
        kp = [("Orçado (BAC)", r_cj.get("orcT")), ("Valor agregado", r_cj.get("medT")), ("Custo lançado (total)", r_cj.get("custoT")),
              ("Estoque financeiro", r_cj.get("estoque")), ("Custo ajustado (itens medidos)", r_cj.get("custoAj")),
              ("Custo em itens sem medição PLS", r_cj.get("custoSem")), ("IDC ajustado", r_cj.get("idc")),
              ("Resultado", r_cj.get("resultado")), ("Estimativa ao término", r_cj.get("eac")), ("Sobra (+) / estouro (−)", r_cj.get("eacDif"))]
        for k, (rot, v) in enumerate(kp, start=tr + 4):
            wr.cell(k, 4, rot)
            wr.cell(k, 5, v).number_format = moeda if rot != "IDC ajustado" else "0.00"

        # ---- Contratos
        wc = wb.create_sheet("%s Contratos" % nome)
        cabecalho(wc, 1, ["Origem (CC)", "Item", "Descrição do item", "Contrato", "Fornecedor", "Tipo", "Descrição do contrato",
                          "Valor contratado", "Lançado no item", "Adiantado", "% lançado do contratado"],
                  [12, 9, 34, 11, 38, 13, 50, 16, 16, 16, 13])
        for c in cj.get("contratos", []):
            wc.append([c.get("cc"), c.get("item"), c.get("nmItem"), c.get("ct"), c.get("forn"), c.get("tipo"), c.get("desc"),
                       c.get("total"), c.get("lancado"), c.get("adiantado"), None])
            r = wc.max_row
            wc.cell(r, 11, "=IF(H%d>0,I%d/H%d,0)" % (r, r, r)).number_format = pct
            for k in (8, 9, 10):
                wc.cell(r, k).number_format = moeda
        n_c = wc.max_row
        wc.cell(n_c + 2, 7, "TOTAL").font = neg
        for k in (9, 10):
            cel = wc.cell(n_c + 2, k, "=SUM(%s2:%s%d)" % (L(k), L(k), max(n_c, 2)))
            cel.number_format, cel.font = moeda, neg

        # ---- Contratos a medir e OC a receber (base da estimativa comprometida)
        wm = wb.create_sheet("%s Contratos (valor total)" % nome)
        cabecalho(wm, 1, ["Origem (CC)", "Item", "Descrição do item", "Contrato", "Fornecedor", "Descrição do contrato", "Valor total", "Medido", "A medir"],
                  [12, 9, 34, 11, 38, 50, 16, 16, 16])
        for c in cj.get("contratosAm", []):
            wm.append([c.get("cc"), c.get("item"), c.get("nmItem"), c.get("ct"), c.get("forn"), c.get("desc"), c.get("total"), c.get("medido"), c.get("saldo")])
            for k in (7, 8, 9):
                wm.cell(wm.max_row, k).number_format = moeda
        n_m = wm.max_row
        wm.cell(n_m + 2, 6, "TOTAL").font = neg
        for k in (7, 8, 9):
            cel = wm.cell(n_m + 2, k, "=SUM(%s2:%s%d)" % (L(k), L(k), max(n_m, 2)))
            cel.number_format, cel.font = moeda, neg
        wa = wb.create_sheet("%s OCs (valor total)" % nome)
        cabecalho(wa, 1, ["Origem (CC)", "Item", "Descrição do item", "OC", "Credor", "Valor da OC", "% da OC neste item", "Valor neste item"],
                  [12, 9, 34, 9, 40, 16, 12, 16])
        for o in cj.get("ocsAr", []):
            wa.append([o.get("cc"), o.get("item"), o.get("nmItem"), o.get("oc"), o.get("credor"), o.get("valor"), o.get("parte"), o.get("valorItem")])
            r = wa.max_row
            for k in (6, 8):
                wa.cell(r, k).number_format = moeda
            wa.cell(r, 7).number_format = pct
        n_a = wa.max_row
        wa.cell(n_a + 2, 5, "TOTAL").font = neg
        cel = wa.cell(n_a + 2, 8, "=SUM(H2:H%d)" % max(n_a, 2))
        cel.number_format, cel.font = moeda, neg

        # ---- Ordens de compra
        wo = wb.create_sheet("%s OCs" % nome)
        cabecalho(wo, 1, ["Origem (CC)", "Item", "Descrição do item", "OC", "Credor", "Lançado", "NF recebida?", "Itens comprados"],
                  [12, 9, 34, 9, 40, 16, 12, 80])
        for o in cj.get("ocs", []):
            wo.append([o.get("cc"), o.get("item"), o.get("nmItem"), o.get("oc"), o.get("credor"), o.get("lancado"),
                       "SIM" if o.get("nf") else "NÃO", o.get("itens")])
            wo.cell(wo.max_row, 6).number_format = moeda
        n_o = wo.max_row
        wo.cell(n_o + 2, 5, "TOTAL").font = neg
        cel = wo.cell(n_o + 2, 6, "=SUM(F2:F%d)" % max(n_o, 2))
        cel.number_format, cel.font = moeda, neg
    bio = io.BytesIO()
    wb.save(bio)
    bio.seek(0)
    return bio, "Auditoria_PLS_%s.xlsx" % time.strftime("%Y%m%d_%H%M")


def auditoria_titulos_xlsx(payload, usuario=""):
    """Excel de auditoria: todos os títulos de pagamento + conferência 'contratos (valor total) + OCs' x lançado.
    payload = {centros:[{cc,nome,orcado}], titulos:[...], contratos:[...], pedidos:[...]} (montado pela tela)."""
    import io
    import openpyxl
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter as L

    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    cab_fill = PatternFill("solid", fgColor="A50E24")
    cab_font = Font(bold=True, color="FFFFFF")
    sub_fill = PatternFill("solid", fgColor="F3E8E8")
    neg = Font(bold=True)
    moeda = '#,##0.00;[Red]-#,##0.00'

    def cab(ws, linha, titulos, larguras, congela=1):
        for i, t in enumerate(titulos, start=1):
            c = ws.cell(linha, i, t)
            c.font, c.fill = cab_font, cab_fill
            c.alignment = Alignment(wrap_text=True, vertical="center", horizontal="center")
        ws.row_dimensions[linha].height = 32
        for i, w in enumerate(larguras, start=1):
            ws.column_dimensions[L(i)].width = w
        ws.freeze_panes = "%s%d" % (L(congela + 1), linha + 1)
        ws.auto_filter.ref = "A%d:%s%d" % (linha, L(len(titulos)), linha)

    titulos = payload.get("titulos", [])
    contratos = payload.get("contratos", [])
    pedidos = payload.get("pedidos", [])
    centros = payload.get("centros", [])

    # ---- Títulos (todos os pagamentos) -- primeiro para as fórmulas das outras abas
    wt = wb.create_sheet("Títulos")
    cab(wt, 1, ["CC", "Item da EAP", "Descrição do item", "Credor", "Título", "Documento", "Grupo", "Contrato", "OC", "Descrição / observação",
                "Status", "Data de pagamento", "Vencimento", "Valor do título", "Pago", "A pagar"],
        [11, 9, 34, 40, 12, 14, 11, 11, 9, 50, 10, 14, 14, 16, 16, 16], congela=4)
    for t in titulos:
        wt.append([t.get("cc"), t.get("item"), t.get("nmItem"), t.get("credor"), t.get("numtit"), t.get("documento"), t.get("grupo"),
                   t.get("contrato"), t.get("oc"), t.get("descr"), t.get("status"), t.get("pagamento"), t.get("vencimento"),
                   t.get("vlr_ap"), t.get("vlr_pg"), t.get("vlr_a_pagar")])
    n_t = max(wt.max_row, 2)
    for r in range(2, n_t + 1):
        for c in (14, 15, 16):
            wt.cell(r, c).number_format = moeda
    wt.cell(n_t + 2, 13, "TOTAL").font = neg
    for c in (14, 15, 16):
        cel = wt.cell(n_t + 2, c, "=SUBTOTAL(109,%s2:%s%d)" % (L(c), L(c), n_t))
        cel.number_format, cel.font, cel.fill = moeda, neg, sub_fill

    def T(col):
        return "'Títulos'!$%s$2:$%s$%d" % (col, col, n_t)

    # ---- Contratos
    wc = wb.create_sheet("Contratos")
    cab(wc, 1, ["CC", "Contrato", "Fornecedor", "Valor total do contrato", "Valor medido", "Saldo a medir", "Lançado em títulos (fórmula)",
                "Medido − lançado", "Títulos (qtd)"],
        [11, 11, 44, 18, 18, 18, 20, 18, 10], congela=3)
    for c in contratos:
        wc.append([c.get("cc"), c.get("ct"), c.get("forn"), c.get("total"), c.get("medido"), c.get("saldo")])
        r = wc.max_row
        num = str(c.get("ct", "")).replace("CTS/", "")
        wc.cell(r, 7, "=SUMIFS(%s,%s,A%d,%s,\"CTS/%s\")" % (T("N"), T("A"), r, T("H"), num))
        wc.cell(r, 8, "=E%d-G%d" % (r, r))
        wc.cell(r, 9, "=COUNTIFS(%s,A%d,%s,\"CTS/%s\")" % (T("A"), r, T("H"), num))
        for k in (4, 5, 6, 7, 8):
            wc.cell(r, k).number_format = moeda
    n_c = max(wc.max_row, 2)
    wc.cell(n_c + 2, 3, "TOTAL").font = neg
    for k in (4, 5, 6, 7, 8):
        cel = wc.cell(n_c + 2, k, "=SUBTOTAL(109,%s2:%s%d)" % (L(k), L(k), n_c))
        cel.number_format, cel.font, cel.fill = moeda, neg, sub_fill

    # ---- OCs
    wo = wb.create_sheet("OCs")
    cab(wo, 1, ["CC", "OC", "Credor", "Classificação", "Data da OC", "Valor da OC", "Realizado + a pagar", "Saldo da OC", "Lançado em títulos (fórmula)",
                "Títulos (qtd)", "Valor da OC neste CC"],
        [11, 9, 44, 30, 13, 18, 18, 18, 20, 10, 20], congela=3)
    for o in pedidos:
        wo.append([o.get("cc"), o.get("oc"), o.get("credor"), o.get("classe"), o.get("data"), o.get("valor"), o.get("realizado"), o.get("saldo"), None, None, o.get("valorCC")])
        r = wo.max_row
        wo.cell(r, 9, "=SUMIFS(%s,%s,A%d,%s,B%d)" % (T("N"), T("A"), r, T("I"), r))
        wo.cell(r, 10, "=COUNTIFS(%s,A%d,%s,B%d)" % (T("A"), r, T("I"), r))
        for k in (6, 7, 8, 9, 11):
            wo.cell(r, k).number_format = moeda
    n_o = max(wo.max_row, 2)
    wo.cell(n_o + 2, 5, "TOTAL").font = neg
    for k in (6, 7, 8, 9, 11):
        cel = wo.cell(n_o + 2, k, "=SUBTOTAL(109,%s2:%s%d)" % (L(k), L(k), n_o))
        cel.number_format, cel.font, cel.fill = moeda, neg, sub_fill

    # ---- Por credor
    wk = wb.create_sheet("Por credor")
    cab(wk, 1, ["Credor", "Títulos (qtd)", "Valor lançado", "Pago", "A pagar", "Primeiro pagamento", "Último pagamento"],
        [52, 12, 18, 18, 18, 16, 16], congela=1)
    credores = {}
    for t in titulos:
        k = (t.get("credor") or "").strip() or "(sem credor)"
        d = credores.setdefault(k, {"n": 0, "mins": [], "ult": []})
        d["n"] += 1
        if t.get("pagamentoIso"):
            d["mins"].append(t["pagamentoIso"])
    for i, k in enumerate(sorted(credores, key=lambda x: x.upper()), start=2):
        wk.cell(i, 1, k)
        wk.cell(i, 2, "=COUNTIFS(%s,A%d)" % (T("D"), i))
        wk.cell(i, 3, "=SUMIFS(%s,%s,A%d)" % (T("N"), T("D"), i))
        wk.cell(i, 4, "=SUMIFS(%s,%s,A%d)" % (T("O"), T("D"), i))
        wk.cell(i, 5, "=SUMIFS(%s,%s,A%d)" % (T("P"), T("D"), i))
        ds = sorted(credores[k]["mins"])
        if ds:
            wk.cell(i, 6, "%s/%s/%s" % (ds[0][8:10], ds[0][5:7], ds[0][:4]))
            wk.cell(i, 7, "%s/%s/%s" % (ds[-1][8:10], ds[-1][5:7], ds[-1][:4]))
        for c in (3, 4, 5):
            wk.cell(i, c).number_format = moeda
    n_k = max(wk.max_row, 2)
    wk.cell(n_k + 2, 1, "TOTAL").font = neg
    for c in (2, 3, 4, 5):
        cel = wk.cell(n_k + 2, c, "=SUBTOTAL(109,%s2:%s%d)" % (L(c), L(c), n_k))
        cel.number_format, cel.font, cel.fill = (moeda if c > 2 else "0"), neg, sub_fill

    # ---- Validação (primeira aba)
    wv = wb.create_sheet("Validação", 0)
    wv.column_dimensions["A"].width = 64
    for i, w in enumerate([20, 20, 20, 20], start=2):
        wv.column_dimensions[L(i)].width = w
    wv.cell(1, 1, "Validação — lançado × contratos (valor total) + OCs").font = Font(bold=True, size=14)
    f_u = payload.get("filtroUau") or {}
    wv.cell(2, 1, "Gerado em %s por %s. Filtro de data do UAU na tela: %s a %s (vazio = sem limite). Todas as linhas abaixo são fórmulas sobre as abas Títulos, Contratos e OCs." % (
        time.strftime("%d/%m/%Y %H:%M"), usuario or "-", f_u.get("ini") or "início", f_u.get("fim") or "hoje"))
    colunas = [c.get("cc") for c in centros]
    for j, c in enumerate(centros, start=2):
        cel = wv.cell(4, j, "%s\n%s" % (c.get("cc"), c.get("nome") or ""))
        cel.font, cel.fill = cab_font, cab_fill
        cel.alignment = Alignment(wrap_text=True, horizontal="center", vertical="center")
    wv.cell(4, 1, "Item").font = cab_font
    wv.cell(4, 1).fill = cab_fill
    wv.row_dimensions[4].height = 34
    tot_col = 2 + len(centros)
    cel = wv.cell(4, tot_col, "Total")
    cel.font, cel.fill = cab_font, cab_fill
    cel.alignment = Alignment(horizontal="center", vertical="center")
    wv.column_dimensions[L(tot_col)].width = 20

    def soma_cc(aba, col_val, col_cc, extra=""):
        def f(letra):
            return "=SUMIFS('%s'!$%s$2:$%s$2000,'%s'!$%s$2:$%s$2000,%s$4%s)" % (aba, col_val, col_val, aba, col_cc, col_cc, letra, extra)
        return f

    def ref(aba, col, n):
        return "'%s'!$%s$2:$%s$%d" % (aba, col, col, n)

    def sumifs(aba, col_val, n, extra=""):
        # soma por CC (o cabeçalho tem "CC\nnome": usamos a linha 3 oculta com o código)
        return lambda letra: "=SUMIFS(%s,%s,%s$3%s)" % (ref(aba, col_val, n), ref(aba, "A", n), letra, extra)

    for j, cc in enumerate(colunas, start=2):
        wv.cell(3, j, cc).font = Font(color="FFFFFF", size=1)
    wv.row_dimensions[3].hidden = True
    wv.cell(3, 1, "(código do CC)")

    linhas = [
        ("A. LANÇADO NOS TÍTULOS (o que o UAU já registrou)", None, None, True),
        ("Orçado do centro de custo", lambda L_, j: centros[j - 2].get("orcado"), None, False),
        ("Títulos lançados — contratos", sumifs("Títulos", "N", n_t, ',%s,"Contrato"' % ref("Títulos", "G", n_t)), None, False),
        ("Títulos lançados — OCs", sumifs("Títulos", "N", n_t, ',%s,"OC"' % ref("Títulos", "G", n_t)), None, False),
        ("Títulos lançados — avulsos (sem contrato e sem OC)", sumifs("Títulos", "N", n_t, ',%s,"Avulso"' % ref("Títulos", "G", n_t)), None, False),
        ("Total lançado em títulos", "SOMA3", None, False),
        ("   dos quais já pagos", sumifs("Títulos", "O", n_t), None, False),
        ("   dos quais a pagar", sumifs("Títulos", "P", n_t), None, False),
        ("B. CONTRATOS (valor total, independente do medido)", None, None, True),
        ("Valor total dos contratos", sumifs("Contratos", "D", n_c), None, False),
        ("   valor medido", sumifs("Contratos", "E", n_c), None, False),
        ("   saldo a medir", sumifs("Contratos", "F", n_c), None, False),
        ("   lançado em títulos nos contratos", sumifs("Contratos", "G", n_c), None, False),
        ("   diferença medido − lançado (medido ainda sem título)", "CT_DIF", None, False),
        ("C. ORDENS DE COMPRA", None, None, True),
        ("OCs próprias (sem contrato) — valor total", sumifs("OCs", "K", n_o, ',%s,"OC própria"' % ref("OCs", "D", n_o)), None, False),
        ("OCs sem título ainda — valor total", sumifs("OCs", "K", n_o, ',%s,"OC sem título"' % ref("OCs", "D", n_o)), None, False),
        ("OCs ligadas a contrato — valor total (já dentro do contrato)", sumifs("OCs", "K", n_o, ',%s,"OC de contrato"' % ref("OCs", "D", n_o)), None, False),
        ("   saldo das OCs próprias + sem título", "OC_SALDO", None, False),
        ("   lançado em títulos nas OCs próprias", sumifs("OCs", "I", n_o, ',%s,"OC própria"' % ref("OCs", "D", n_o)), None, False),
        ("   diferença valor da OC − lançado (OC ainda sem título)", "OC_DIF", None, False),
        ("D. COMPROMETIDO — definição: contratos (valor total) + OCs (valor total) + avulsos", None, None, True),
        ("Comprometido", "COMPROM", None, False),
        ("Orçado − comprometido (sobra +, estouro −)", "RESULT", None, False),
        ("E. CONFERÊNCIA COM O PAINEL", None, None, True),
        ("Lançado calculado pelo painel (item a item)", lambda L_, j: centros[j - 2].get("painel"), None, False),
        ("Diferença painel − comprometido (deve ser 0,00)", "DIFPAINEL", None, False),
    ]
    pos = {}
    r = 5
    for rot, fn, _, secao in linhas:
        wv.cell(r, 1, rot)
        if secao:
            for c in range(1, tot_col + 1):
                wv.cell(r, c).fill = sub_fill
            wv.cell(r, 1).font = neg
        pos[rot] = r
        r += 1
    P = pos
    for rot, fn, _, secao in linhas:
        if secao:
            continue
        row = P[rot]
        for j in range(2, tot_col):
            col = L(j)
            if callable(fn):
                try:
                    v = fn(col)
                except TypeError:
                    v = fn(col, j)
                wv.cell(row, j, v)
            elif fn == "SOMA3":
                wv.cell(row, j, "=%s%d+%s%d+%s%d" % (col, P["Títulos lançados — contratos"], col, P["Títulos lançados — OCs"], col, P["Títulos lançados — avulsos (sem contrato e sem OC)"]))
            elif fn == "CT_DIF":
                wv.cell(row, j, "=%s%d-%s%d" % (col, P["   valor medido"], col, P["   lançado em títulos nos contratos"]))
            elif fn == "OC_SALDO":
                wv.cell(row, j, "=SUMIFS(%s,%s,%s$3,%s,\"<>OC de contrato\")" % (ref("OCs", "H", n_o), ref("OCs", "A", n_o), col, ref("OCs", "D", n_o)))
            elif fn == "OC_DIF":
                wv.cell(row, j, "=%s%d-%s%d" % (col, P["OCs próprias (sem contrato) — valor total"], col, P["   lançado em títulos nas OCs próprias"]))
            elif fn == "COMPROM":
                wv.cell(row, j, "=%s%d+%s%d+%s%d+%s%d" % (col, P["Valor total dos contratos"], col, P["OCs próprias (sem contrato) — valor total"],
                                                          col, P["OCs sem título ainda — valor total"], col, P["Títulos lançados — avulsos (sem contrato e sem OC)"]))
            elif fn == "RESULT":
                wv.cell(row, j, "=%s%d-%s%d" % (col, P["Orçado do centro de custo"], col, P["Comprometido"]))
            elif fn == "DIFPAINEL":
                wv.cell(row, j, "=%s%d-%s%d" % (col, P["Lançado calculado pelo painel (item a item)"], col, P["Comprometido"]))
        # total
        wv.cell(row, tot_col, "=SUM(B%d:%s%d)" % (row, L(tot_col - 1), row))
        for j in range(2, tot_col + 1):
            wv.cell(row, j).number_format = moeda
        if rot in ("Total lançado em títulos", "Comprometido", "Lançado calculado pelo painel (item a item)", "Orçado − comprometido (sobra +, estouro −)"):
            for j in range(1, tot_col + 1):
                wv.cell(row, j).font = neg
    r = max(P.values()) + 2
    notas = [
        "Leitura: 'Comprometido' (= lançado do painel) conta cada contrato pelo valor total (mesmo o que ainda não foi medido) e cada OC pelo valor total, mais os títulos avulsos.",
        "OCs ligadas a contrato não somam de novo: o valor delas já está dentro do contrato. 'Lançado em títulos' é só conferência do que o UAU já registrou.",
        "Datas: 'Data de pagamento' é a baixa do título (só para pagos); 'Vencimento' é a data prevista dos títulos a pagar.",
    ]
    for k, tx in enumerate(notas):
        wv.cell(r + k, 1, tx)
    wv.freeze_panes = "B5"
    bio = io.BytesIO()
    wb.save(bio)
    bio.seek(0)
    return bio, "Auditoria_Titulos_%s.xlsx" % time.strftime("%Y%m%d_%H%M")
